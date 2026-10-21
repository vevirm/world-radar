"""Reader-first Word briefing for the living foresight atlas.

The report is a map of the collected research and curated changes, not a
prediction test. All possible follow-up links are retained in a readable
reading trail; the complete 1,366-record collection stays in the atlas.
"""
from __future__ import annotations
from collections import defaultdict
from pathlib import Path


def write_report(atlas, report):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.shared import Inches, Pt, RGBColor

    # One monochrome theme for both local builds and future scheduled GitHub runs.
    BLACK = RGBColor(0, 0, 0)
    NAVY = TEAL = MUTED = CORAL = BLACK
    doc = Document()
    sec=doc.sections[0]
    sec.top_margin=Inches(.78);sec.bottom_margin=Inches(.74)
    sec.left_margin=Inches(.78);sec.right_margin=Inches(.78)
    normal=doc.styles['Normal']
    normal.font.name='Arial';normal.font.size=Pt(10.4);normal.font.color.rgb=BLACK
    normal.paragraph_format.space_after=Pt(6)
    normal.paragraph_format.line_spacing=1.19
    for key,size in [('Title',26),('Heading 1',16.5),('Heading 2',12.5)]:
        sty=doc.styles[key]
        sty.font.name='Arial';sty.font.size=Pt(size);sty.font.color.rgb=BLACK;sty.font.bold=True
        sty.paragraph_format.space_before=Pt(16)
        sty.paragraph_format.space_after=Pt(8)
        sty.paragraph_format.keep_with_next=True
    h=sec.header.paragraphs[0]
    h.text='WORLD RADAR  /  THE LIVING FUTURES ATLAS'
    h.runs[0].font.size=Pt(8);h.runs[0].font.color.rgb=TEAL;h.runs[0].bold=True
    f=sec.footer.paragraphs[0]
    f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    run=f.add_run('Living research briefing   •   ')
    run.font.size=Pt(8);run.font.color.rgb=MUTED
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)

    def p(txt='', lead=None, small=False):
        para=doc.add_paragraph()
        if lead:
            r=para.add_run(lead);r.bold=True;r.font.color.rgb=TEAL
        # Bold numeric findings without changing the source text or guessing at outcomes.
        import re
        emphasized = bool(lead and lead.strip() in ('WHAT CAME NEXT', 'WHERE IT IS HEADING', 'STILL RELEVANT?', 'STILL OPEN'))
        for piece in re.split(r'(?<!\w)(\d{1,3}(?:,\d{3})+|\d+(?:\.\d+)?%)(?!\w)', str(txt)):
            if not piece:
                continue
            r=para.add_run(piece)
            if re.fullmatch(r'\d{1,3}(?:,\d{3})+|\d+(?:\.\d+)?%', piece):
                r.bold=True
            if small:r.font.size=Pt(9.2);r.font.color.rgb=BLACK
        if emphasized:
            para.paragraph_format.space_after=Pt(8)
        return para

    def hyperlink(par, value, url):
        if not str(url).startswith('https://'): return
        rid=par.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
        node=OxmlElement('w:hyperlink');node.set(qn('r:id'),rid)
        r=OxmlElement('w:r');pr=OxmlElement('w:rPr')
        col=OxmlElement('w:color');col.set(qn('w:val'),'000000');pr.append(col)
        und=OxmlElement('w:u');und.set(qn('w:val'),'single');pr.append(und)
        r.append(pr);txt=OxmlElement('w:t');txt.text=value.replace(' ↗','');r.append(txt);node.append(r);par._p.append(node)

    def label(text):
        para=doc.add_paragraph()
        r=para.add_run(text.upper())
        r.font.size=Pt(8.5);r.font.color.rgb=BLACK;r.bold=True
        para.paragraph_format.space_before=Pt(13)
        para.paragraph_format.space_after=Pt(4)
        para.paragraph_format.keep_with_next=True
        return para

    def sources(items, original=None):
        para=doc.add_paragraph()
        para.paragraph_format.space_after=Pt(7)
        if original: hyperlink(para,'Original ↗',original)
        for ix,ref in enumerate(items[:5],1):
            if ix>1 or original:para.add_run('   ·   ')
            hyperlink(para, f'Source {ix} ↗', ref.get('url',''))

    summary=atlas['summary']
    label('A tour of ideas that are still moving')
    doc.add_heading('Back then.\nWhere are we now?', 0)
    opening=p('What Europe studied, hoped, planned and worried about — and how the conversation has changed. '
            'Some things have happened. Some are taking shape. Some were never meant to be predictions. '
            'This edition keeps those differences visible.')
    opening.runs[0].font.size=Pt(13);opening.runs[0].font.color.rgb=MUTED
    p(f"{summary['total']:,} original research records  •  {summary['compared_with_literature']:,} checked against later scholarly literature  •  "
      f"{summary['with_later_reading']} with possible later reading  •  {summary['editorial_contexts']} source-checked contemporary snapshots")
    doc.add_heading('The idea in five questions',1)
    for lead,detail in (
      ('What did they say then?  ', 'Keep the original context: forecasts, scenarios, choices and warning signs are different things.'),
      ('What do we know now?  ', 'Use observations when available, and label later academic discussion as discussion.'),
      ('Where does it seem to be heading?  ', 'Describe direction and conditions, not an imaginary final result.'),
      ('Is it still relevant?  ', 'Sometimes yes, sometimes transformed, sometimes unreviewed — never assume an old question has expired.'),
      ('Can it be rated?  ', 'A 1–100 editorial assessment is optional. It needs a clear object, a reason, and comparable evidence.'),
    ):p(detail,lead=lead)
    p('This is a curated reading guide to a complete searchable research atlas. It is not a completed verification of all 1,366 entries.',small=True)

    doc.add_heading('What has actually changed?',1)
    p('Four contemporary public-policy threads, with dated references. They provide wider context — they are not presented as things the indexed academic papers necessarily predicted.')
    for item in atlas.get('editorial_context',[]):
        label(item['tag'])
        doc.add_heading(item['title'],2)
        p(item['then'],lead='BACK THEN  ')
        p(item['now'],lead='WHAT CAME NEXT  ')
        p(item['direction'],lead='WHERE IT IS HEADING  ')
        p(item['relevance'],lead='STILL RELEVANT?  ')
        p(item['open_question'],lead='STILL OPEN  ')
        sources(item.get('sources',[]))

    doc.add_heading('Eight conversations across time',1)
    p('This part looks across the whole collection, grouped by subject rather than by whether an idea can be scored. '
      'The examples are entry points into the archive, not a claim that the collection has a single unified conclusion.')
    by_theme=defaultdict(list)
    for row in atlas['records']:by_theme[row['topic']].append(row)
    for t in atlas['themes']:
        doc.add_heading(t['headline'],2)
        p(t['intro'])
        p(f"In the archive: {t['count']} research items; {t['later_reading']} with candidate later scholarship. "
          "This is a coverage count, not a relevance or quality score.",small=True)
        featured=[r for r in by_theme[t['topic']] if r['followups']]
        featured=sorted(featured,key=lambda r:(0 if r['followups'][0].get('abstract_excerpt') else 1,r['year'],r['title']))[:4]
        for rec in featured:
            label(f"{rec['year'] or 'Undated'} / an idea that kept being discussed")
            p(rec['title'],lead='THEN  ')
            first=rec['followups'][0]
            p(first['title']+ (f" ({first['year']})" if first.get('year') else ''),lead='LATER READING  ')
            sources([first],original=rec['original_url'])
        if not featured:p('No credible potential later reading link has yet been catalogued for this theme.')
        p('WHAT IS STILL OPEN  The presence of a later paper does not establish a real-world trajectory. Checking present-day outcomes and usefulness is the next editorial question.',small=True)

    doc.add_heading('Reading trail: every linked follow-up',1)
    p(f"The research scan found {summary['with_later_reading']} older entries with possibly related later publications. "
      'All are listed below, grouped by subject, so none vanishes behind a handful of highlight cards. '
      'These are search leads, not a claim of direct intellectual descent or proof that the original future happened.')
    for theme in atlas['themes']:
        rows=[r for r in by_theme[theme['topic']] if r['followups']]
        if not rows:continue
        doc.add_heading(f"{theme['topic']} ({len(rows)})",2)
        for row in sorted(rows,key=lambda x:(x['year'],x['title'].lower())):
            para=doc.add_paragraph()
            para.paragraph_format.space_after=Pt(9)
            para.paragraph_format.keep_together=True
            r=para.add_run(f"{row['year'] or '?'}  {row['title']}  ")
            r.bold=True;r.font.size=Pt(9)
            para.add_run('\n')
            hyperlink(para,'Original ↗',row['original_url'])
            for later in row['followups']:
                para.add_run('   ·   ')
                hyperlink(para,f"{later['year'] or '?'} later reading ↗",later['url'])

    doc.add_heading('How to read a hindsight score',1)
    p('If a topic can be fairly evaluated, the point is to judge the thinking — clarity of assumptions, usefulness for planning and how the expected direction compares with evidence. '
      'A scenario is not wrong merely because one branch did not happen; a 2030 target is not a forecast.')
    p('Not yet rated means not reviewed or not comparable, not zero. Percent-of-target figures are descriptive progress measures, not quality scores. '
      'Direction, changed relevance and evidence strength should appear even where a numerical score would be misleading.')
    doc.add_heading('How complete is this edition?',1)
    p(f"Every one of the {summary['total']:,} indexed research records is searchable in the accompanying website and JSON atlas. "
      f"{summary['with_later_reading']} have recorded candidate later scholarly links; "
      f"{summary['with_later_abstract_finding']} have a later abstract containing a result-like passage, not necessarily an observed outcome. "
      'Most entries still lack a source-reviewed real-world comparison. The gaps are visible rather than hidden.')
    p('Next editorial task: review the strongest older ideas topic by topic, add comparable current statistics and counterevidence, '
      'then write a short, dated "where we are heading" and "is it still relevant" assessment for each. Never populate those fields with guesses.')

    # Harmonize the full document, including Word's hyperlink XML and header/footer.
    for para in doc.paragraphs:
        if para.style.name == 'Heading 1':
            ppr = para._p.get_or_add_pPr()
            border = OxmlElement('w:pBdr')
            rule = OxmlElement('w:bottom')
            rule.set(qn('w:val'), 'single')
            rule.set(qn('w:sz'), '4')
            rule.set(qn('w:color'), '000000')
            border.append(rule)
            ppr.append(border)
        for run in para.runs:
            run.font.color.rgb = BLACK
            if para.text.startswith('In the archive:') and any(ch.isdigit() for ch in run.text):
                run.bold = True
            if para.text.startswith('WHAT IS STILL OPEN') and run.text.startswith('WHAT IS STILL OPEN'):
                run.bold = True
    for part in (sec.header, sec.footer):
        for para in part.paragraphs:
            for run in para.runs:
                run.font.color.rgb = BLACK
    # python-docx's default Title style includes a blue border; turn it black too.
    # This also prevents colored theme accents in later Word template versions.
    for root in (doc.styles.element, doc._element, sec.header._element, sec.footer._element):
        for node in root.xpath('.//w:color'):
            node.set(qn('w:val'), '000000')
            for attr in ('themeColor', 'themeShade', 'themeTint'):
                node.attrib.pop(qn('w:' + attr), None)
        for node in root.xpath('.//w:bottom'):
            if node.get(qn('w:color')) is not None:
                node.set(qn('w:color'), '000000')
                node.attrib.pop(qn('w:themeColor'), None)
    report=Path(report);report.parent.mkdir(parents=True,exist_ok=True);doc.save(report)
    return report
