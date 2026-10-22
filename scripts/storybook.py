"""Human-readable hindsight stories: audited observations, not an accuracy leaderboard.

A literary follow-up is a *reading suggestion*, not proof that a historic
prediction came true. Numerical hindsight ratings require source-reviewed
original claims and independent comparable observed outcomes.
"""
from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
import json
import re

ROOT = Path(__file__).resolve().parents[1]
STORIES = ROOT / 'hindsight-stories.json'


def clean(value, limit=None):
    txt = re.sub(r'\s+', ' ', str(value or '')).strip()
    return txt[:limit-1].rstrip() + '…' if limit and len(txt) > limit else txt


def hindsight_rating(case):
    """Compute an editor's 1–100 hindsight grade only with an actual review.

    Clarity, anticipation, and usefulness are *judgments*, not probabilities.
    Each must be documented and accompanied by a reason; a mere source hit
    cannot be scored. Valid cases have already passed outcome_verification.
    """
    review = case.get('hindsight_review') or {}
    dimensions = ('clarity', 'anticipation', 'usefulness')
    if not review or not clean(review.get('reason')):
        return None
    if not all(type(review.get(key)) is int and 1 <= review[key] <= 100 for key in dimensions):
        return None
    value = int(.25*review["clarity"] + .5*review["anticipation"] + .25*review["usefulness"] + .5)
    return {'value': value, 'label': 'Editor-reviewed hindsight rating',
            'reason': clean(review['reason'], 260),
            'rubric': {'clarity': review['clarity'], 'anticipation': review['anticipation'],
                      'usefulness': review['usefulness']}}


def progress_meter(item):
    """Separate goal-proximity yardstick, NEVER a forecast-quality score."""
    indicator = item.get('progress_indicator') or {}
    try:
        value, goal = float(indicator['value']), float(indicator['goal'])
    except (KeyError, ValueError, TypeError):
        return None
    if value < 0 or goal <= 0 or not clean(indicator.get('metric')) or not clean(indicator.get('as_of')):
        return None
    return {'value': min(100, round(100 * value / goal)), 'label': 'Target-level yardstick (NOT a hindsight grade)',
            'metric': clean(indicator['metric']), 'as_of': clean(indicator['as_of']),
            'calculation': f'{value:g} ÷ {goal:g} × 100',
            'note': clean(indicator.get('note')) or 'Shows the measured level relative to a goal, not forecast accuracy.'}


def make_stories(entries, reviewed, observations):
    by_id = {e['id']: e for e in entries}
    stories = []
    for number, item in enumerate(observations):
        # Historical policy snapshots remain in outcome-observations.json, but
        # only a snapshot explicitly associated with a CURRENT Radar 1 original
        # can appear in Radar 2's then/now output.
        marker = clean(item.get('original_lead_title_contains')).casefold()
        if not marker or not any(marker in clean(e.get('title')).casefold() for e in entries):
            continue
        sources = [s for s in item.get('sources', []) if s.get('url', '').startswith('https://')]
        # An observation must provide subsequent sourced results to enter the front page.
        if not sources or not item.get('expected') or not item.get('happened'):
            continue
        stories.append({'id': f'policy-{number}', 'kind': 'policy_progress',
            'tag': 'A promise in progress', 'title': clean(item.get('topic'), 125),
            'then': clean(item['expected'], 580), 'now': clean(item['happened'], 580),
            'take': clean(item.get('verdict'), 300), 'hindsight_rating': None,
            'progress_meter': progress_meter(item), 'sources': sources[:5]})
    for number, case in enumerate(reviewed):
        lead = by_id.get(case.get('lead_id'), {})
        original = case['original']
        obs = case['observation']
        stories.append({'id': f'forecast-{number}', 'kind': 'reviewed_forecast',
            'tag': 'Reviewed prediction', 'title': clean(lead.get('title') or 'Historic prediction', 125),
            'then': clean(original['verbatim_passage'], 560),
            'now': clean(obs['what_happened'], 580),
            'take': clean(case['verdict'].replace('_', ' ').capitalize() + '. ' + case.get('limitations', ''), 320),
            'hindsight_rating': hindsight_rating(case), 'progress_meter': None,
            'sources': [{'url': original['url'], 'date': original['published_on'], 'finding': 'Original passage'}]
                + obs['sources'][:4]})
    # Only use a small, readable set of *provisional* scholarly follow-ups.
    # These are deliberately not placed in the actual-outcomes list.
    # Human-friendly, carefully limited paraphrases of the three pinned papers.
    # The title match is necessary but NOT an assertion the older idea was right.
    wanted = (
        ('Record dry summer in 2015 challenges precipitation projections in Central Europe',
         'In 2016, a study questioned how confidently climate projections describe Central Europe’s summers.',
         'A related 2018 modeling paper found very different possible summer-rainfall futures: wet, dry and very dry.',
         'Varying soil moisture'),
        ('Assessment of Future Heat Stress Risk in European Regions',
         'A 2017 study explored how changing climates and social conditions might affect Europe’s heat-stress risks.',
         'A related 2021 study modeled rising heat-related mortality risks through 2050, with markedly different projections across European countries.',
         'Future Projections of Heat Mortality Risk'),
        ('By 2050 the Mitigation Effects of EU Forests Could Nearly Double',
         'A 2017 paper explored whether climate-smart forestry could nearly double forests’ climate-mitigation contribution by 2050.',
         'A related 2022 study examined forest-management trade-offs, including the role of unmanaged and broad-leaved forests in carbon storage and local cooling.',
         'Trade‐Offs for Climate‐Smart Forestry'),
    )
    shelf = []
    for title_fragment, plain_then, plain_later, needed_followup_title in wanted:
        item = next((e for e in entries if title_fragment.lower() in clean(e.get('title')).lower()
                     and e.get('follow_up_publications')), None)
        if not item:
            continue
        studies = [s for s in item['follow_up_publications'] if s.get('url', '').startswith('https://')]
        if not studies:
            continue
        chosen = next((s for s in studies if needed_followup_title.casefold() in clean(s.get('title')).casefold()), None)
        if not chosen:
            continue  # Don't attach a prepared summary to a different study.
        abstract = clean(chosen.get('finding_from_abstract'), 310)
        shelf.append({'id': item['id'], 'title': clean(item['title'], 110),
            'year': clean(item.get('original_publication'))[:4],
            'then': plain_then,
            'later': plain_later if abstract else f'A later study followed up: {clean(chosen.get("title"), 160)}.',
            'take': 'This is later academic discussion, not a real-world test or proof that the earlier scenario happened.',
            'original_url': item['source_url'],
            'later_url': chosen['url'], 'later_title': clean(chosen.get('title'), 135)})
    return {'generated_utc': datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC'),
        'summary': {'research_leads': len(entries), 'reviewed_forecasts': len(reviewed),
                    'policy_snapshots': sum(x['kind'] == 'policy_progress' for x in stories),
                    'later_reading': len(shelf)},
        'stories': stories, 'reading_shelf': shelf,
        'scoring_note': ('Only an original-source-reviewed prediction with comparable independent outcome evidence '
                         'can receive a hindsight grade. The 1–100 rating is editorial: clarity 25%, '
                         'anticipation 50%, usefulness 25%. Policy target-level meters are separate and '
                         'do not measure prediction quality. No evidence means no grade.')}


def write_stories(entries, reviewed, observations, path=None):
    data = make_stories(entries, reviewed, observations)
    target = Path(path) if path else STORIES
    target.write_text(json.dumps(data, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    return data


def write_word(stories, report, *, run=None):
    """A short magazine-like Word narrative; research internals remain in JSON."""
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.style import WD_STYLE_TYPE

    NAVY = RGBColor(27, 52, 64)
    TEAL = RGBColor(21, 104, 102)
    GREY = RGBColor(88, 107, 114)
    CORAL = RGBColor(182, 86, 60)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(.82)
    section.bottom_margin = Inches(.75)
    section.left_margin = Inches(.92)
    section.right_margin = Inches(.92)
    section.header_distance = Inches(.38)
    section.footer_distance = Inches(.35)
    normal = doc.styles['Normal']
    normal.font.name, normal.font.size, normal.font.color.rgb = 'Aptos', Pt(10.5), NAVY
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.16
    for name in ('Title', 'Heading 1', 'Heading 2'):
        s = doc.styles[name]
        s.font.name = 'Aptos Display'
        s.font.color.rgb = NAVY
        s.paragraph_format.space_before = Pt(12)
        s.paragraph_format.space_after = Pt(6)
        s.paragraph_format.keep_with_next = True
    doc.styles['Title'].font.size = Pt(33)
    doc.styles['Heading 1'].font.size = Pt(19)
    doc.styles['Heading 2'].font.size = Pt(13)
    if 'Eyebrow' not in doc.styles:
        eyebrow = doc.styles.add_style('Eyebrow', WD_STYLE_TYPE.PARAGRAPH)
        eyebrow.font.name = 'Aptos'
        eyebrow.font.size = Pt(8)
        eyebrow.font.bold = True
        eyebrow.font.color.rgb = TEAL
        eyebrow.paragraph_format.space_after = Pt(5)
    else:
        eyebrow = doc.styles['Eyebrow']
    header = section.header.paragraphs[0]
    header.text = 'THE HINDSIGHT FILES  /  EUROPEAN FORESIGHT RADAR'
    header.style = doc.styles['Eyebrow']
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    f = footer.add_run('Research digest  •  ')
    f.font.size = Pt(8)
    f.font.color.rgb = GREY
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    footer._p.append(fld)

    def label(text):
        p = doc.add_paragraph(style='Eyebrow')
        p.add_run(text.upper())
        return p

    def paragraph(text, *, bold_prefix=None):
        p = doc.add_paragraph()
        if bold_prefix:
            r = p.add_run(bold_prefix)
            r.bold = True
            p.add_run(clean(text))
        else:
            p.add_run(clean(text))
        return p

    def link(p, text, uri):
        # Clickable hyperlinks rather than long unwieldy URL walls.
        from docx.opc.constants import RELATIONSHIP_TYPE
        from docx.oxml import OxmlElement
        if not uri or not str(uri).startswith('https://'):
            return
        rel = p.part.relate_to(uri, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
        h = OxmlElement('w:hyperlink')
        h.set(qn('r:id'), rel)
        run = OxmlElement('w:r')
        prop = OxmlElement('w:rPr')
        color = OxmlElement('w:color'); color.set(qn('w:val'), '156866'); prop.append(color)
        under = OxmlElement('w:u'); under.set(qn('w:val'), 'single'); prop.append(under)
        run.append(prop)
        t = OxmlElement('w:t'); t.text = text
        run.append(t); h.append(run); p._p.append(h)

    label('Historical European futures, revisited')
    doc.add_heading('Back then.\nAnd now?', 0)
    p = paragraph('What Europeans hoped, planned or expected — and what the evidence says happened next. Short stories, not 100 pages of audit logs.')
    p.runs[0].font.size = Pt(14)
    p.runs[0].font.color.rgb = GREY
    doc.add_paragraph()
    summary = stories['summary']
    t = doc.add_table(rows=1, cols=3)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, number, name in zip(t.rows[0].cells, (summary['policy_snapshots'], summary['reviewed_forecasts'], summary['research_leads']),
                                  ('Real-world snapshots', 'Reviewed predictions', 'Research leads behind the scenes')):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'EAF3F0'); cell._tc.get_or_add_tcPr().append(shade)
        pp = cell.paragraphs[0]
        pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rr = pp.add_run(str(number) + '\n'); rr.bold = True; rr.font.size = Pt(21); rr.font.color.rgb = TEAL
        rr2 = pp.add_run(name); rr2.font.size = Pt(8); rr2.font.color.rgb = NAVY
    doc.add_paragraph()
    doc.add_heading('The idea', 1)
    paragraph('Each story has three simple beats: what was being said, what came afterward, and a short verdict. We do not pretend every scenario was a prediction, or that an academic paper proves something happened in the real world.')
    doc.add_heading('About the numbers', 1)
    paragraph('A score out of 100 can be useful — when it has earned its place. Hindsight grades require an editor to check the original prediction against later independent measurements. For a policy target, we can show a separate 1–100 target-level yardstick; it is not a score for how clever the original idea was.')
    p = paragraph('Only source-reviewed comparisons will be presented as outcomes. This is a research-stage edition until those comparisons are available.' if not stories['stories'] else 'These examples use saved, dated source material. They are snapshots, not claims about the situation today.')
    p.runs[0].italic = True
    if stories['stories']:
        heading = doc.add_heading('What later evidence actually shows', 1)
        heading.paragraph_format.page_break_before = True
        paragraph('Every comparison shown below comes from a current accepted historical original with a documented later evidence trail. Policy targets are not predictions.')
    else:
        doc.add_heading('Research status: evidence still missing',1)
        paragraph('The accepted historical collection is ready, but no original has yet been paired with a source-reviewed, comparable later real-world development. We will not fill this gap with unrelated academic links or treat an imagined scenario as a forecast.')
        paragraph('The next tasks are to extract passages from curated EU institutional reports, follow the actual questions from accepted originals, and check later measurements and implementation against them.')
    for index, story in enumerate(stories['stories']):
        if index:
            doc.add_paragraph('')
        label('01 / policy progress' if index == 0 else ('02 / policy progress' if index == 1 else f'{index+1:02d} / reviewed forecast'))
        doc.add_heading(story['title'], 2)
        paragraph(story['then'], bold_prefix='BACK THEN  ')
        paragraph(story['now'], bold_prefix='WHAT CAME NEXT  ')
        verdict = doc.add_paragraph()
        r = verdict.add_run('THE TAKE  ')
        r.bold = True
        r.font.color.rgb = CORAL
        verdict.add_run(story['take'])
        if story.get('hindsight_rating'):
            grade = story['hindsight_rating']
            p = paragraph(f"{grade['value']}/100 — {grade['reason']}", bold_prefix='HINDSIGHT GRADE  ')
            p.runs[0].font.color.rgb = TEAL
        elif story.get('progress_meter'):
            meter = story['progress_meter']
            p = paragraph(f"{meter['value']}/100 target-level yardstick ({meter['calculation']}). {meter['metric']}; data: {meter['as_of']}. {meter['note']}", bold_prefix='PROGRESS SNAPSHOT  ')
            p.runs[0].font.color.rgb = TEAL
        else:
            p = paragraph('Not rated — no source-reviewed like-for-like outcome comparison yet.', bold_prefix='HINDSIGHT GRADE  ')
            p.runs[0].font.color.rgb = TEAL
        pp = doc.add_paragraph()
        pp.paragraph_format.space_after = Pt(5)
        rr = pp.add_run('Explore the evidence   ')
        rr.font.size = Pt(8.5)
        rr.font.bold = True
        rr.font.color.rgb = GREY
        for i, s in enumerate(story.get('sources', [])[:4], 1):
            link(pp, f"{clean(s.get('date') or 'Source')} source {i}", s.get('url'))
            pp.add_run('   ')

    if stories['reading_shelf']:
        heading = doc.add_heading('Meanwhile, in the reading room…', 1)
        heading.paragraph_format.page_break_before = True
        paragraph('These are interesting conversations across time. A later paper can deepen or challenge an older idea, but that is not the same as independently checking whether the world turned out as expected.')
        for index, item in enumerate(stories['reading_shelf'][:2], 1):
            label(f'{index:02d} / research echo, not an outcome')
            doc.add_heading(item['title'], 2)
            paragraph(item['then'], bold_prefix='BACK THEN  ')
            paragraph(item['later'], bold_prefix='LATER RESEARCH  ')
            p = paragraph(item['take'], bold_prefix='VERDICT  ')
            p.runs[0].font.color.rgb = CORAL
            sources = doc.add_paragraph()
            link(sources, 'Original study', item['original_url'])
            sources.add_run('    ·    ')
            link(sources, 'Later study', item['later_url'])

    heading = doc.add_heading('How to read a hindsight score', 1)
    heading.paragraph_format.page_break_before = bool(stories['stories'] or stories['reading_shelf'])
    paragraph('The future is messy. So should the judging be curious rather than overconfident.')
    label('A score is editorial, not a probability')
    for title, body in (
        ('25%  Clarity', 'Could someone at the time tell what was expected, when, and for whom?'),
        ('50%  Anticipation', 'How closely did the expectation match later measured developments, including important misses?'),
        ('25%  Usefulness', 'Did it offer a concrete way to plan, prepare or make a decision?'),
    ):
        paragraph(body, bold_prefix=title + '   ')
    doc.add_heading('When there is no number', 2)
    paragraph('That is usually a good thing: it means the original forecast has not yet been checked in its full source, the relevant date has not arrived, or the later evidence is about a different place, metric or modeled scenario. A missing score is not zero.')
    doc.add_heading('For the curious', 2)
    paragraph(f"There are {summary['research_leads']:,} indexed research leads behind this short digest, and {summary['reviewed_forecasts']} reviewed, comparable forecast outcomes. They remain available in the project’s data files. This Word file is meant for reading, not for auditing the scanner.")
    run = run or {}
    if run.get('provider'):
        paragraph('Latest research mode: ' + ('Saved catalogue only — no new web requests' if run['provider'] in ('catalogue', 'evidence')
                  else 'Automated research-assisted follow-up — its results still need source review') + '.')
    paragraph('No prediction is graded just because a related academic abstract exists. The goal is an enjoyable, traceable retrospective — not an impressive-looking but unjustified accuracy table.')
    Path(report).parent.mkdir(parents=True, exist_ok=True)
    doc.save(report)
