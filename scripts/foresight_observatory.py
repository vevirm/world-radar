#!/usr/bin/env python3
"""Cumulative, *provisional* claim-to-source research. No automated verdicts.

Runs after Radar 1; exports a public claim catalogue and a separate human/LLM
review queue. A title or shared general field is never outcome evidence.
"""
from __future__ import annotations
import argparse, datetime as dt, hashlib, json, os, re, time, urllib.parse, urllib.request
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt

ROOT=Path(__file__).resolve().parents[1]
STORE='observatory-data.json'
CLAIMS='observatory-claims.json'
QUEUE='observatory-possible-connections.json'
REVIEW='observatory-reviewed.json'
STOP=set(('study research studies report analysis impact future futures foresight scenarios forecast forecasting prediction projected projection prospective european europe union eu the and for from into on of in at with by to than vs over under between across review approaches effects assessment based modelling model models evidence results towards changing change changes possible potential opportunities challenges comparative case method methods supply demand towards level perspective developments development').split())
DOMAIN_GROUPS={
 'aviation':('aviation','aircraft','airline','airlines','jet','saf','sustainable aviation fuel','biojet'),
 'flood':('flood','flooding','inundation','floodplain','flood damage'),
 'carbon':('carbon','emissions trading','emission trading','ets','carbon pricing','allowance','cap-and-trade'),
 'hydrogen':('hydrogen','fuel cell','h2'),
 'electric vehicle':('electric vehicle','ev adoption','battery electric','electric car'),
 'renewable electricity':('renewable electricity','solar','wind power','renewable power'),
}
EU_WORDS=re.compile(r'\b(europe|european|eu|eurostat|member states|european union|european commission|finland|france|germany|italy|spain|sweden|netherlands|poland|belgium|austria|ireland|portugal|denmark|norway|greece|czech|romania|hungary|baltic)\b',re.I)
TOKEN=re.compile(r'[a-z][a-z0-9-]{2,}')

def words(s):return [x for x in TOKEN.findall(str(s or '').lower()) if x not in STOP and len(x)>3]
def stamp():return dt.datetime.now(dt.timezone.utc).isoformat(timespec='seconds')
def getjson(p,default):
 try:return json.loads(Path(p).read_text(encoding='utf8'))
 except FileNotFoundError:return default

def source_claims(root=ROOT):
 data=getjson(root/'research-leads.json',{'leads':[]})['leads']; result=[]
 for row in data:
  date=str(row.get('date') or ''); year=int(date[:4]) if date[:4].isdigit() else 0
  if not (2011<=year<=2022):continue
  if row.get('radar1_selection') not in ('active',None):continue
  seed=row.get('radar2_seed') or {}
  passage=str(seed.get('original_passage') or row.get('original_passage') or '').strip()
  if len(passage)<65:continue
  if not (row.get('eu_connection') or EU_WORDS.search(row.get('title','')+' '+passage)):continue
  doi=str(row.get('doi') or '').lower().removeprefix('https://doi.org/')
  id='claim-'+hashlib.sha256((doi or row.get('url','') or row.get('title','')).encode()).hexdigest()[:16]
  result.append({'id':id,'original_id':doi,'title':row.get('title',''),'year':year,'date':date,'original_claim':passage,'original_url':row.get('url',''),'doi':doi,'status':'original_claim_unreviewed','last_seen':row.get('last_seen','')})
 return result

def subjects(claim):
 text=claim['title']+' '+claim['original_claim']; lo=text.lower()
 group=[key for key,forms in DOMAIN_GROUPS.items() if any(re.search(r'(?<!\w)'+re.escape(term)+r'(?!\w)',lo) for term in forms)]
 # Keep distinctive object words from title, not generic shared concepts.
 candidates=list(dict.fromkeys(words(claim['title'])))
 return group, candidates[:8]

def queries(claim):
 groups, tokens=subjects(claim)
 terms=[]
 for group in groups[:2]:terms.append(group)
 terms += tokens
 terms=list(dict.fromkeys(terms))
 if len(terms)<2:return []
 geographic=' Europe' if EU_WORDS.search(claim['title']+' '+claim['original_claim']) else ''
 q=[' '.join(terms[:3])+geographic]
 if groups:q.append(groups[0]+' '+ ' '.join(x for x in tokens if x not in groups[0].split())[:90]+geographic)
 return list(dict.fromkeys(x.strip() for x in q if len(x.split())>=2))[:2]

def relationship(claim,item):
 title=str(item.get('title') or ''); abstract=str(item.get('abstract') or ''); content=(title+' '+abstract).lower()
 if not title or not abstract or len(abstract)<80:return None
 year=int(str(item.get('year') or '0')[:4]);
 if year<=claim['year'] or year>dt.date.today().year:return None
 groups,tokens=subjects(claim); discovered=[g for g in groups if any(re.search(r'(?<!\w)'+re.escape(term)+r'(?!\w)',content) for term in DOMAIN_GROUPS[g])]
 if groups and not discovered:return None  # reject cereal supply vs aviation
 match=[w for w in tokens if w in set(words(title+' '+abstract))]
 if not groups and len(match)<3:return None
 if groups and len(match)<2:return None
 scope='EU/European named in source' if EU_WORDS.search(content) else 'European relevance not established'
 # Out-of-region paper may be explanatory background but not direct outcome evidence.
 if scope!='EU/European named in source' and not re.search(r'\b(global|international|worldwide)\b',content):return None
 kind='candidate_direct_or_explanatory' if scope.startswith('EU/') else 'possible_background_only'
 return {'relationship':kind,'match_basis':'specific object: '+(', '.join(discovered) or ', '.join(match[:3]))+'; shared claim concepts: '+', '.join(match[:5]),'geographic_scope':scope,'evidence_status':'metadata_only_not_read_or_checked'}

def inverted(idx):
 if not isinstance(idx,dict):return ''
 out=['']*(min(3500,max((max(v) for v in idx.values() if v),default=-1)+1))
 for word,positions in idx.items():
  for p in positions:
   if 0<=p<len(out):out[p]=word
 return ' '.join(out)

def search_openalex(query,year,api_key='',timeout=18):
 params={'search':query,'filter':f'from_publication_date:{year+1}-01-01','per_page':20,
         'select':'id,doi,title,publication_year,publication_date,abstract_inverted_index,primary_location'}
 if api_key:params['api_key']=api_key
 url='https://api.openalex.org/works?'+urllib.parse.urlencode(params)
 req=urllib.request.Request(url,headers={'User-Agent':'EuropeanForesightObservatory/1.0 (research discovery)','Accept':'application/json'})
 with urllib.request.urlopen(req,timeout=timeout) as response:payload=json.load(response)
 for r in payload.get('results',[]):
  loc=r.get('primary_location') or {}
  yield {'title':r.get('title') or '', 'year':r.get('publication_year'), 'date':r.get('publication_date'),
    'url':r.get('doi') or (loc.get('landing_page_url') if loc else '') or r.get('id'),
    'abstract':inverted(r.get('abstract_inverted_index') or {}),'index':'OpenAlex'}

def make_docs(root,db):
 queue=db.get('connections',{}); claims={c['id']:c for c in db['claims']}; items=[x for arr in queue.values() for x in arr]
 def init(title,subtitle):
  d=Document();s=d.sections[0];s.header_distance=Inches(.4);s.footer_distance=Inches(.4)
  d.add_heading(title,0);d.add_paragraph(subtitle);d.add_paragraph('Generated: '+stamp()[:10])
  return d
 d=init('Possible connections — UNCHECKED','Machine-selected candidates, not evidence that a foresight claim came true or failed. Source abstracts are index metadata; full texts have not been verified.')
 d.add_paragraph(f'{len(claims)} historical claims in library · {len(items)} provisional claim–source links. Search is cumulative; unfilled claims remain open.')
 for cid in sorted(queue):
  if not queue[cid] or cid not in claims:continue
  c=claims[cid];d.add_heading(f"{c['year']} · {c['title']}",level=2);d.add_paragraph('Original passage: '+c['original_claim']);d.add_paragraph('Original: '+c['original_url'])
  for x in queue[cid][:8]:
   d.add_heading(x['title'],level=3)
   for line in ('Published: '+str(x['year']), 'Source: '+x['url'], 'Why candidate: '+x['match_basis'], 'Scope: '+x['geographic_scope'], 'Status: UNCHECKED; metadata only', 'Abstract extract: '+x['abstract'][:950]):d.add_paragraph(line)
 if not items:d.add_paragraph('No strong provisional matches saved yet. This is not evidence that nothing happened.')
 d.save(root/'POSSIBLE-CONNECTIONS-UNCHECKED.docx')
 approved=getjson(root/REVIEW,{'findings':[]})
 d=init('European futures — reviewed findings','Only human/LLM-reviewed findings supplied in observatory-reviewed.json appear here. No scanner-generated conclusion is published as reviewed.')
 if not approved.get('findings'):d.add_paragraph('No reviewed developmental interpretations yet. Add approved entries to observatory-reviewed.json; this Word file will update on the next export.')
 for x in approved.get('findings',[]):
  if x.get('review_status')!='approved':continue
  d.add_heading(x.get('title') or x.get('claim_id','Review'),level=2)
  d.add_paragraph('Claim: '+x.get('original_claim',''))
  d.add_paragraph('Interpretation: '+x.get('interpretation',''))
  d.add_paragraph('Direction: '+x.get('direction','unclassified'))
  for source in x.get('sources',[]):d.add_paragraph('Evidence: '+str(source))
 d.save(root/'EU-FORESIGHT-REVIEWED.docx')

def export(root,db):
 db['updated_at']=stamp(); (root/STORE).write_text(json.dumps(db,ensure_ascii=False,indent=2)+'\n',encoding='utf8')
 (root/CLAIMS).write_text(json.dumps({'updated_at':db['updated_at'],'claims':db['claims']},ensure_ascii=False,indent=2)+'\n',encoding='utf8')
 (root/QUEUE).write_text(json.dumps({'updated_at':db['updated_at'],'connections':db['connections']},ensure_ascii=False,indent=2)+'\n',encoding='utf8')
 make_docs(root,db)

def run(root=ROOT,limit=4,max_seconds=600,searcher=search_openalex):
 root=Path(root);existing=getjson(root/STORE,{'schema_version':1,'claims':[],'connections':{},'investigations':{},'updated_at':''})
 byid={c['id']:c for c in existing['claims']}
 for c in source_claims(root):byid[c['id']]=c
 existing['claims']=sorted(byid.values(),key=lambda c:(-c['year'],c['title']))
 investigations=existing.setdefault('investigations',{});connections=existing.setdefault('connections',{})
 candidates=sorted(existing['claims'],key=lambda c:(investigations.get(c['id'],{}).get('attempts',0),investigations.get(c['id'],{}).get('last_run',''),c['id']))[:limit]
 started=time.monotonic();errors=[];found=0
 for c in candidates:
  cid=c['id'];old=investigations.get(cid,{});record={'attempts':old.get('attempts',0)+1,'last_run':stamp(),'query_terms':queries(c),'error':None}
  investigations[cid]=record;seen={x['url'] for x in connections.get(cid,[])}
  for q in queries(c):
   if time.monotonic()-started>max_seconds:break
   try:
    for paper in searcher(q,c['year'],os.getenv('OPENALEX_API_KEY','')):
     if not paper.get('url') or paper['url'] in seen:continue
     rel=relationship(c,paper)
     if not rel:continue
     seen.add(paper['url']);connections.setdefault(cid,[]).append({**paper,**rel,'claim_id':cid,'found_at':stamp(),'query':q});found+=1
   except Exception as e:
    record['error']=type(e).__name__+': '+str(e)[:180];errors.append({'claim_id':cid,'error':record['error']})
  # Save after each claim, including API failures. No conclusion or 'nothing happened' from empty results.
  export(root,existing)
 export(root,existing)
 return {'investigated':len(candidates),'new_candidates':found,'errors':errors,'total_claims':len(existing['claims'])}

if __name__=='__main__':
 p=argparse.ArgumentParser();p.add_argument('--limit',type=int,default=4);p.add_argument('--seconds',type=int,default=600);p.add_argument('--export-only',action='store_true');a=p.parse_args()
 if a.export_only:
  d=getjson(ROOT/STORE,{'schema_version':1,'claims':source_claims(ROOT),'connections':{},'investigations':{}});export(ROOT,d);print('Exported documents and web JSON')
 else:print(json.dumps(run(limit=a.limit,max_seconds=a.seconds),indent=2))
