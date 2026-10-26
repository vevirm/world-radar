#!/usr/bin/env python3
"""EU Futures Observatory: cumulative, provisional claim-to-literature discovery.
Never interprets a candidate source as evidence of a realised outcome.
"""
import argparse, hashlib, json, os, re, time, urllib.parse, urllib.request, urllib.error, zipfile
from datetime import datetime, timezone
from pathlib import Path
from xml.sax.saxutils import escape
from connection_groups import read_decisions, suggest_groups

ROOT=Path(__file__).resolve().parents[1]
STOP=set('a an and are as at be been being by can could do for from had has have how if in into is it its may might more most of on or our should some such than that the their these this through to toward towards under using was were what when where which while will with would year years study studies research report paper results evidence analysis future futures forecasting forecast scenarios scenario projection projected long term eu european europe'.split())
SYNONYMS={
 'aviation':['aviation','aircraft','airline','jet','flight','saf'],
 'fuel':['fuel','biofuel','kerosene','e-fuel','saf','sustainable aviation'],
 'carbon':['carbon','co2','emissions','emission','greenhouse','ghg'],
 'pricing':['pricing','price','prices','trading','ets','allowance','tax'],
 'flood':['flood','flooding','inundation'],
 'energy':['energy','electricity','power'],
 'hydrogen':['hydrogen','h2'],
 'demographic':['demographic','population','ageing','aging'],
 'employment':['employment','labour','labor','jobs','workforce'],
 'renewable':['renewable','solar','wind','clean energy'],
}
EU_RE=re.compile(r'\b(europe|european|eu|eurostat|eu-\d+|european union|\b(?:finland|germany|france|italy|spain|sweden|poland|netherlands|belgium|portugal|ireland|greece|austria|denmark|norway|estonia|latvia|lithuania|czechia|slovakia|slovenia|croatia|romania|bulgaria|hungary|cyprus|malta|luxembourg)\b)\b',re.I)

def words(s):return re.findall(r'[a-z][a-z0-9-]{2,}',(s or '').lower())
def keywords(s,limit=12):
 result=[]
 for w in words(s):
  if w not in STOP and not w.isnumeric() and w not in result:result.append(w)
 return result[:limit]
def jsonread(path,default):
 try:return json.loads(Path(path).read_text(encoding='utf8'))
 except (FileNotFoundError,ValueError):return default

# Editorial corrections validated against the historical catalogue, 2026-10-24.
# Applied after a stable claim ID is derived so pre-existing links remain intact.
EDITORIAL_YEARS = {
 '10.1080/21699763.2017.1411288': 2018,
 '10.2478/jlecol-2018-0020': 2018,
 '10.1515/eec-2017-0004': 2017,
 '10.1017/s1867299x00000593': 2011,
}
EDITORIAL_ORDER_DOI = '10.7163/gpol.0008'
EDITORIAL_ORDER_TEXT = (
 'For this purpose two alternative reference scenarios are outlined and examined in the light of some specific characteristics of Poland’s metropolitan development. '
 'The questions posed in this context involve the sustainability of policy assumptions concerning the role of cities of various size categories, as well as factors that could lead to a discontinuity in the trends observed in metropolisation.'
)

def apply_editorial_corrections(claim):
 doi=(claim.get('doi') or '').strip().lower().removeprefix('https://doi.org/')
 if doi in EDITORIAL_YEARS:
  year=EDITORIAL_YEARS[doi]
  claim['year']=year
  claim['date']=f'{year}-01-01'  # year-only correction; exact issue date not asserted
 if doi==EDITORIAL_ORDER_DOI:
  claim['original_claim']=EDITORIAL_ORDER_TEXT
 claim['original_claim']=re.sub(r'^Abstract\s+', '',claim.get('original_claim') or '')
 return claim

def claims():
 data=jsonread(ROOT/'research-leads.json',{}).get('leads',[])
 output=[]
 for x in data:
  date=(x.get('date') or '')[:10]
  if not ('2011-01-01'<=date<'2023-01-01'):continue
  if x.get('radar1_selection') not in ('active',None):continue
  seed=x.get('radar2_seed') or {}; passage=(seed.get('original_passage') or '').strip()
  if len(passage)<45:continue
  # Original excerpts from abstracts remain explicitly labelled as provisional excerpts.
  identity=x.get('doi') or x.get('url') or x.get('title')
  ident='claim-'+hashlib.sha256((identity+'|'+passage).encode()).hexdigest()[:16]
  output.append({'id':ident,'title':x.get('title',''),'year':int(date[:4]),'date':date,
    'original_claim':passage,'source_url':x.get('url',''),'doi':x.get('doi',''),
    'source_status':seed.get('source_status','unverified'),'theme':x.get('foresight_scope','') or '',
    'source_field':seed.get('source_field','unknown'),'review_status':'source excerpt not independently verified'})
 return [apply_editorial_corrections(x) for x in output]

def anchors(claim):
 raw=claim['original_claim']+' '+claim['title']
 terms=keywords(raw,32)
 # less weight to title; must have genuine anchor terms from original proposition
 base=keywords(claim['original_claim'],18)
 return list(dict.fromkeys(base+terms))[:18]

def queries(claim):
 a=anchors(claim)
 if not a:return []
 # use short, differentiated semantically meaningful fragments, not title prefixes.
 group=a[:4]; extra=[x for x in a[4:] if x not in group]
 region='Europe' if EU_RE.search(claim['original_claim']+' '+claim['title']) else ''
 one=' '.join(group[:3]+([region] if region else []))
 two=' '.join((group[0:2]+extra[0:2])+([region] if region else []))
 return list(dict.fromkeys([one.strip(),two.strip()]))

def concept_sets(s):
 tokens=set(words(s)); sets=[]
 for name,terms in SYNONYMS.items():
  if name in tokens or any(t in s.lower() for t in terms):
   sets.append((name,terms))
 return sets

def compare(claim,paper):
 text=((paper.get('title') or '')+' '+(paper.get('abstract') or '')).lower()
 if not text.strip():return None
 original=(claim['original_claim']+' '+claim['title']).lower()
 if paper.get('year',0)<=claim['year']:return None
 claim_terms=keywords(claim['original_claim'],25)
 paper_terms=set(words(text))
 hits=[w for w in claim_terms if w in paper_terms]
 groups=concept_sets(original)
 group_hits=[name for name,alts in groups if any(re.search(r'(?<!\w)'+re.escape(alt)+r'(?!\w)',text) for alt in alts)]
 # The central object must match: a generic word like costs/supply alone is insufficient.
 generic={'cost','costs','supply','demand','prices','price','policy','economic','growth','change','changes','development','impact','effects','risk','risks','global','future','market','markets','model','models','new','social','public'}
 distinctive=[w for w in hits if w not in generic]
 if len(distinctive)<2 and not (group_hits and len(distinctive)>=1):return None
 if len(hits)<3 and len(group_hits)<2:return None
 claim_eu=bool(EU_RE.search(original)); paper_eu=bool(EU_RE.search(text))
 # Require European scope for a strong link; global/missing geography is review-only.
 rank='strong_candidate' if len(distinctive)>=3 and (not claim_eu or paper_eu) else 'possible_candidate'
 if claim_eu and not paper_eu:rank='possible_candidate'
 return {'rank':rank,'matched_terms':distinctive[:12],'concepts':group_hits,
   'geography_overlap':('eu_in_later_text' if paper_eu else 'eu_not_established' if claim_eu else 'not_required'),
   'reason':'Shared claim-specific concepts: '+', '.join(distinctive[:6])+'; '+('EU scope visible' if paper_eu else 'EU scope not established')+'. This is a provisional connection, not an outcome finding.'}

def fetch_openalex(q,after_year,limit=20,page=1):
 params={'search':q,'filter':f'from_publication_date:{after_year+1}-01-01,to_publication_date:2026-10-23','per-page':str(limit),'page':str(page),'select':'id,doi,title,display_name,publication_year,primary_location,abstract_inverted_index'}
 url='https://api.openalex.org/works?'+urllib.parse.urlencode(params)
 headers={'User-Agent':'EU-Futures-Observatory/1.0 (research; public sources)'}
 api_key=os.environ.get('OPENALEX_API_KEY','').strip()
 if not api_key:
  raise RuntimeError('OPENALEX_API_KEY is missing from the Scanner 2 job environment; verify the GitHub Actions secret mapping')
 headers['Authorization']='Bearer '+api_key
 req=urllib.request.Request(url,headers=headers)
 with urllib.request.urlopen(req,timeout=25) as response:raw=json.load(response)
 results=[]
 for p in raw.get('results',[]):
  inv=p.get('abstract_inverted_index') or {}; ordered=[]
  for token,locs in inv.items():
   for i in locs:ordered.append((i,token))
  abstract=' '.join(t for _,t in sorted(ordered))[:6000]
  loc=p.get('primary_location') or {}
  results.append({'id':p.get('id',''),'title':p.get('title') or p.get('display_name') or '',
      'year':p.get('publication_year') or 0,'doi':p.get('doi') or '',
      'url':p.get('doi') or p.get('id') or '', 'abstract':abstract,
      'landing_url':loc.get('landing_page_url') or '','source':'OpenAlex'})
 return results

def scan(minutes=20,max_claims=400,max_papers=20,min_minutes=15):
 start=time.monotonic(); pool=claims(); store=jsonread(ROOT/'observatory-data.json',{'schema_version':1,'claims':{},'connections':{},'history':[],'cursor':0})
 store['claims'].update({x['id']:x for x in pool})
 for _claim in store['claims'].values():
  apply_editorial_corrections(_claim)
 if not pool:raise RuntimeError('No eligible historical claim excerpts')
 cursor=int(store.get('cursor') or 0)%len(pool)
 # Run through further result pages instead of stopping after the first shallow pass.
 # Persist the next page for every claim/query so later runs do not fetch the same page again.
 progress=store.setdefault('search_progress',{})
 processed=set(); errors=[]; tried=0; retained=0; searches=0; failed_this_run=set()
 budget_end=start+minutes*60-35
 target_end=start+min_minutes*60
 stop_reason='budget_exhausted'
 # Continue from pages already investigated on earlier runs. A larger bounded
 # horizon leaves further evidence to explore while rotating across claims.
 max_page=30
 request_interval=3.0  # cap API traffic instead of exhausting/rate-limiting in minutes
 next_request_at=start
 rate_limit_strikes=0
 while time.monotonic()<budget_end:
  work_found=False
  # A round-robin sweep ensures the run explores many claims before going deeper.
  for _ in range(min(len(pool),max_claims)):
   if time.monotonic()>=budget_end:break
   claim=pool[cursor];cursor=(cursor+1)%len(pool)
   available=[]
   for q in queries(claim):
    key=claim['id']+'|'+q
    page=int(progress.get(key,1))
    if page<=max_page and key not in failed_this_run:available.append((key,q,page))
   if not available:continue
   # Query with the least advancement first, so both variants are explored.
   key,q,page=min(available,key=lambda item:item[2])
   work_found=True;processed.add(claim['id'])
   known={c.get('source_id') for c in store['connections'].values() if c.get('claim_id')==claim['id']}
   try:
    # Respect the public API: remaining runtime is used for paced research,
    # never for an idle delay added after the search has ended.
    pause=max(0.0,next_request_at-time.monotonic())
    if pause:time.sleep(min(pause,max(0.0,budget_end-time.monotonic())))
    if time.monotonic()>=budget_end:break
    next_request_at=time.monotonic()+request_interval
    candidates=fetch_openalex(q,claim['year'],max_papers,page=page)
    rate_limit_strikes=0
    searches+=1;tried+=len(candidates)
    progress[key]=max_page+1 if len(candidates)<max_papers else page+1
    for paper in candidates:
     if not paper.get('id') or paper['id'] in known:continue
     rel=compare(claim,paper)
     if rel is None:continue
     ident='link-'+hashlib.sha256((claim['id']+'|'+paper['id']).encode()).hexdigest()[:18]
     store['connections'][ident]={'id':ident,'claim_id':claim['id'],'source_id':paper['id'],'source_title':paper['title'],
       'source_year':paper['year'],'source_url':paper['url'],'source_doi':paper['doi'],'abstract':paper['abstract'],
       'retrieved_by':q,'relevance':rel,'review_status':'unchecked','editorial_note':'',
       'collected_at':datetime.now(timezone.utc).isoformat()}
     known.add(paper['id']);retained+=1
   except Exception as e:
    if isinstance(e,urllib.error.HTTPError) and e.code==429:
     rate_limit_strikes+=1
     # Exponential backoff with a ceiling; do not continually hammer OpenAlex.
     next_request_at=max(next_request_at,time.monotonic()+min(90,15*2**min(rate_limit_strikes-1,3)))
    errors.append({'claim_id':claim['id'],'query':q,'page':page,'error':str(e)[:240]})
    # Avoid repeatedly hitting a failing endpoint in the same run. Retry next run.
    progress[key]=page
    failed_this_run.add(key)
   store['cursor']=cursor
   (ROOT/'observatory-data.json').write_text(json.dumps(store,ensure_ascii=False,indent=2)+'\n',encoding='utf8')
  if not work_found:
   stop_reason='all_search_pages_exhausted';break
  if time.monotonic()>=target_end:
   stop_reason='minimum_work_target_reached';break
  # If API requests fail quickly, avoid an endless tight retry cycle.
  if errors and searches==0:
   stop_reason='all_requests_failed';break
 elapsed=round(time.monotonic()-start,1)
 if elapsed<min_minutes*60:
  print(f'EARLY_STOP: {stop_reason}; active research {elapsed}s; target {min_minutes*60:.0f}s',flush=True)
 store['cursor']=cursor
 store['history'].append({'at':datetime.now(timezone.utc).isoformat(),'processed_claims':len(processed),
  'search_requests':searches,'candidates_seen':tried,'new_connections':retained,'errors':errors[:15],
  'research_seconds':elapsed,'minimum_research_target_seconds':min_minutes*60,'stop_reason':stop_reason})
 store['history']=store['history'][-150:]
 (ROOT/'observatory-data.json').write_text(json.dumps(store,ensure_ascii=False,indent=2)+'\n',encoding='utf8')
 print(json.dumps(store['history'][-1],indent=2))
 return store

def make_docx(path,title,subtitle,records):
 from docx import Document
 from docx.shared import Pt, Inches, RGBColor
 from docx.enum.text import WD_ALIGN_PARAGRAPH
 from docx.enum.style import WD_STYLE_TYPE
 d=Document(); sec=d.sections[0]
 sec.top_margin=Inches(.8);sec.bottom_margin=Inches(.75)
 sec.left_margin=Inches(.92);sec.right_margin=Inches(.92)
 styles=d.styles
 normal=styles['Normal'];normal.font.name='Aptos';normal.font.size=Pt(10)
 normal.font.color.rgb=RGBColor(42,53,65)
 normal.paragraph_format.space_after=Pt(6)
 for level,size in [('Title',22),('Heading 2',12)]:
  st=styles[level];st.font.name='Aptos Display';st.font.size=Pt(size)
  st.font.bold=True;st.font.color.rgb=RGBColor(25,67,87)
  st.paragraph_format.space_before=Pt(14);st.paragraph_format.space_after=Pt(6)
 if 'Source Note' not in [x.name for x in styles]:
  small=styles.add_style('Source Note',WD_STYLE_TYPE.PARAGRAPH)
 else:small=styles['Source Note']
 small.font.name='Aptos';small.font.size=Pt(8.5);small.font.color.rgb=RGBColor(92,108,116)
 small.paragraph_format.space_after=Pt(5)
 hdr=sec.header.paragraphs[0];hdr.text='VV FUTURES  /  EUROPEAN FORESIGHT OBSERVATORY'
 hdr.style='Source Note'
 d.add_heading(title,0)
 intro=d.add_paragraph(subtitle);intro.style='Source Note'
 d.add_paragraph('')
 for number,(head,paragraphs) in enumerate(records,1):
  d.add_heading(f'{number}. {head}',level=2)
  for item in paragraphs:
   txt=str(item)
   if not txt.strip():continue
   p=d.add_paragraph(style='Source Note' if txt.startswith(('Original source', 'Editorial status:', 'Dossier sources:', 'Other web sources', 'Available abstract', 'Group ID:','Review status:')) else 'Normal')
   if ':' in txt and len(txt.split(':',1)[0])<35:
    tag,rest=txt.split(':',1);p.add_run(tag+':').bold=True;p.add_run(rest)
   else:p.add_run(txt)
 foot=sec.footer.paragraphs[0];foot.text='VV Futures · Research dossier · Review source links before citing'
 foot.style='Source Note'
 d.save(path)

def build():
 """Regenerate ONLY the provisional working dossier for Let's see.

 Editorial documents (Reviewed Findings, Historical Claims, and the Excel overview)
 are intentionally never written here. They change only through manual curation.
 """
 store=jsonread(ROOT/'observatory-data.json',{'claims':{},'connections':{},'history':[]})
 store.setdefault('claims',{}).update({x['id']:x for x in claims()})
 (ROOT/'observatory-data.json').write_text(json.dumps(store,ensure_ascii=False,indent=2)+'\n',encoding='utf8')
 claims0=sorted(store.get('claims',{}).values(),key=lambda a:(-a.get('year',0),a.get('title','')))
 links=list(store.get('connections',{}).values())
 links.sort(key=lambda x:(x.get('claim_id',''),-int(x.get('source_year',0) or 0)))
 from collections import defaultdict
 by=defaultdict(list)
 for p in links:by[p.get('claim_id')].append(p)
 decisions=read_decisions(ROOT/'connection-group-decisions.json')
 groups=suggest_groups(store.get('claims',{}),links,decisions)
 possible=[('Group suggestions — screened combinations', ['Exact groups marked rejected by an LLM review are excluded here. Individual source leads further below are NOT proposed two-paper groups and remain available for other combinations.'])]
 for g in groups:
  c=store['claims'].get(g['claim_id'],{})
  paras=['Group ID: '+g['id'],'Original source: '+c.get('source_url',''),
    'Original future claim: '+c.get('original_claim','')]
  for p in g['papers']:
   paras.append('Later source: '+p.get('source_title','')+' — '+p.get('source_url',''))
  paras.append('Review status: '+g['status']+' — possible relationship only')
  possible.append(('Claim + two later sources: '+c.get('title',''),paras))
 possible.append(('Individual research leads — not grouped', ['These are single claim–paper research leads, not claim + paper A + paper B proposals. A paper in a rejected group may still be useful in a different combination.']))
 for c in claims0:
  match=by.get(c['id'],[])
  if not match:continue
  paras=['Original source ('+str(c.get('year',''))+'): '+c.get('source_url',''),
   'Original wording ('+c.get('source_field','')+'; '+c.get('source_status','')+'): '+c.get('original_claim','')]
  for p in match[:12]:
   paras.extend([f"POSSIBLE {p.get('source_year','')}: {p.get('source_title','')} — {p.get('source_url','')}",
    (p.get('relevance') or {}).get('reason',''),
    'Available abstract (not full text): '+((p.get('abstract') or '')[:1400] or 'Unavailable')])
  possible.append((c.get('title',''),paras))
 from research_additions import research_records
 if research_records():
  possible.extend([('Research update · 26 October 2026 — independently check before citing',['12 researched cases. External LLM checks are not editorial verification; these connections are not prediction scores.'])]+research_records())
 make_docx(ROOT/'VV-FUTURES-POSSIBLE-CONNECTIONS.docx','Possible connections — not checked',
  'Automatically collected research leads for “Let’s see”. These are NOT verified outcomes or final findings.',
  possible or [('No candidate connections yet',['No candidate sources retained so far.'])])
 print("Updated Let's see:",len(claims0),'claims,',len(links),'unchecked connections;',
       'main-page Word and Excel left unchanged')

if __name__=='__main__':
 parser=argparse.ArgumentParser();parser.add_argument('command',choices=['scan','build']);parser.add_argument('--minutes',type=float,default=10);parser.add_argument('--claims',type=int,default=400);parser.add_argument('--min-minutes',type=float,default=15)
 args=parser.parse_args()
 if args.command=='scan':scan(args.minutes,args.claims,min_minutes=args.min_minutes);build()
 else:build()
