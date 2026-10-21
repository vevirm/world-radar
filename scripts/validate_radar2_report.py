"""Fail CI if Radar 2's Word artifact is absent, invalid, or lacks key sections."""
import json
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / 'stuff' / 'WHAT_WAS_EXPECTED_WHAT_HAPPENED.docx'
AUDIT = ROOT / 'scanner2-audit.json'


def validate(report=REPORT, audit_path=AUDIT):
    if not report.is_file() or report.stat().st_size < 8000:
        raise ValueError(f'Radar 2 Word file missing or too small: {report}')
    try:
        with ZipFile(report) as archive:
            if archive.testzip() is not None:
                raise ValueError('Radar 2 Word file has damaged ZIP entries')
            for required in ('[Content_Types].xml', 'word/document.xml'):
                if required not in archive.namelist():
                    raise ValueError(f'Radar 2 Word file lacks {required}')
        document = Document(report)
    except (BadZipFile, KeyError, OSError) as exc:
        raise ValueError(f'Radar 2 output is not a valid Word document: {exc}') from exc
    text = '\n'.join(p.text for p in document.paragraphs) + '\n' + '\n'.join(
        c.text for t in document.tables for r in t.rows for c in r.cells)
    required_phrases = (
        'Back then.',
        'WHAT CAME NEXT',
        'Two promises, two reality checks',
        'How to read a hindsight score',
    )
    for phrase in required_phrases:
        if phrase not in text:
            raise ValueError(f'Radar 2 Word file is missing a required section: {phrase}')
    audit = json.loads(audit_path.read_text(encoding='utf-8'))
    if 'run' not in audit or 'reviewed_verdicts' not in audit:
        raise ValueError('Radar 2 audit metadata is incomplete')
    print(f'Validated Radar 2 Word document ({report.stat().st_size:,} bytes); '
          f'{audit["research_leads"]} research leads, '
          f'{audit["reviewed_verdicts"]} source-reviewed verdicts.')
    return audit


if __name__ == '__main__':
    validate()
